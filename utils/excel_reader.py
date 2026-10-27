from pathlib import Path
from openpyxl import load_workbook

class ExcelReader:

    def __init__(self):
        self.file_path = (
            Path(__file__).resolve().parent.parent
            / "test_data"
            / "test_data.xlsx"
        )

    def get_test_data(self, sheet_name):
        workbook = load_workbook(
            filename=self.file_path,
            read_only=True,
            data_only=True
        )

        try:
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)

            headers = next(rows)

            test_data = []

            for row in rows:
                if all(value is None for value in row):
                    continue

                row_data = dict(zip(headers, row))
                test_data.append(row_data)

            return test_data

        finally:
            workbook.close()
